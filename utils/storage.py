"""
utils/storage.py — Lưu trữ dữ liệu: Google Sheets + CSV/Excel fallback
=======================================================================
Ưu tiên:
  1. Google Sheets (nếu có credentials.json)
  2. CSV local (nếu không có credentials hoặc kết nối thất bại)

Xuất báo cáo:
  - export_excel() → file .xlsx với 2 sheet:
      Sheet "Chi tiết": bảng từng bản ghi
      Sheet "Thống kê":  pivot đếm theo nhãn cảm xúc
"""

import csv
import os
import time
import logging
from datetime import datetime
from typing import Optional
import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import config
from utils.emotion_tracker import EmotionRecord

logger = logging.getLogger(__name__)


# Header cho CSV / Google Sheets
SHEET_HEADERS = [
    "STT", "Mức hài lòng", "Thời gian", "Độ ổn định (%)",
    "Confidence TB", "Số frame", "Thời lượng (s)",
    "Rất hài lòng", "Hài lòng", "Bình thường", "Không hài lòng",
]


class StorageManager:
    """
    Quản lý lưu trữ bản ghi cảm xúc.

    Tự động phát hiện Google Sheets credentials và fallback sang CSV nếu cần.
    Dữ liệu cũng được lưu in-memory để phục vụ API và xuất Excel.
    """

    def __init__(self):
        self._records: list[EmotionRecord] = []
        self._use_sheets  = False
        self._gc          = None   # gspread client
        self._worksheet   = None   # gspread worksheet object
        self._csv_path    = None   # Path file CSV local
        self._session_id  = datetime.now().strftime("%Y%m%d_%H%M%S")

        # Tạo thư mục exports nếu chưa có
        os.makedirs(config.EXPORTS_DIR, exist_ok=True)

        # Thử kết nối Google Sheets
        self._init_google_sheets()

        # Nếu không dùng Sheets → dùng CSV
        if not self._use_sheets:
            self._init_csv_fallback()

    # ── Khởi tạo Google Sheets ────────────────────────────────────────────────

    def _init_google_sheets(self):
        """Thử kết nối Google Sheets API."""
        if not os.path.exists(config.CREDENTIALS_PATH):
            logger.info("Không tìm thấy credentials.json → dùng CSV local.")
            return

        try:
            import gspread
            from google.oauth2.service_account import Credentials

            scopes = [
                "https://spreadsheets.google.com/feeds",
                "https://www.googleapis.com/auth/drive",
            ]
            creds = Credentials.from_service_account_file(
                config.CREDENTIALS_PATH, scopes=scopes
            )
            self._gc = gspread.authorize(creds)

            # Mở hoặc tạo spreadsheet
            try:
                sh = self._gc.open(config.GOOGLE_SHEET_NAME)
            except gspread.SpreadsheetNotFound:
                sh = self._gc.create(config.GOOGLE_SHEET_NAME)
                sh.share("", perm_type="anyone", role="writer")
                logger.info(f"Đã tạo Google Sheet mới: {config.GOOGLE_SHEET_NAME}")

            # Lấy hoặc tạo worksheet
            try:
                ws = sh.worksheet(config.GOOGLE_WORKSHEET)
            except gspread.WorksheetNotFound:
                ws = sh.add_worksheet(config.GOOGLE_WORKSHEET, rows=1000, cols=20)
                ws.append_row(SHEET_HEADERS)
                logger.info(f"Đã tạo worksheet: {config.GOOGLE_WORKSHEET}")

            self._worksheet  = ws
            self._use_sheets = True
            logger.info("✅ Kết nối Google Sheets thành công!")

        except ImportError:
            logger.warning("gspread chưa cài (pip install gspread google-auth). Dùng CSV.")
        except Exception as e:
            logger.warning(f"Lỗi kết nối Google Sheets: {e}. Dùng CSV fallback.")

    # ── Khởi tạo CSV fallback ─────────────────────────────────────────────────

    def _init_csv_fallback(self):
        """Tạo file CSV local với header."""
        fname = f"records_{self._session_id}.csv"
        self._csv_path = os.path.join(config.EXPORTS_DIR, fname)
        with open(self._csv_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(SHEET_HEADERS)
        logger.info(f"CSV fallback: {self._csv_path}")

    # ── Ghi bản ghi ──────────────────────────────────────────────────────────

    def write_record(self, record: EmotionRecord):
        """
        Ghi một bản ghi vào Google Sheets hoặc CSV.
        Luôn lưu vào in-memory trước.
        """
        self._records.append(record)

        row = self._record_to_row(record)

        if self._use_sheets:
            try:
                self._worksheet.append_row(row, value_input_option="USER_ENTERED")
                logger.debug(f"Đã ghi vào Sheets: {record.emotion} @ {record.timestamp}")
            except Exception as e:
                logger.error(f"Lỗi ghi Sheets: {e}. Lưu vào CSV.")
                self._fallback_write_csv(row)
        else:
            self._fallback_write_csv(row)

    def _fallback_write_csv(self, row: list):
        """Ghi 1 row vào CSV (dùng khi Sheets lỗi)."""
        if self._csv_path is None:
            self._init_csv_fallback()
        with open(self._csv_path, "a", newline="", encoding="utf-8-sig") as f:
            csv.writer(f).writerow(row)

    def _record_to_row(self, r: EmotionRecord) -> list:
        """Chuyển EmotionRecord thành danh sách giá trị cho 1 row."""
        sat_label = r.satisfaction or config.EMOTION_TO_SATISFACTION_VI.get(r.emotion, r.emotion)
        sat_counts = {level: 0 for level in config.SATISFACTION_LEVELS}
        for raw_em, cnt in r.all_counts.items():
            s = config.EMOTION_TO_SATISFACTION_VI.get(raw_em, "Bình thường")
            sat_counts[s] = sat_counts.get(s, 0) + cnt

        return [
            r.record_no,
            sat_label,
            r.timestamp,
            r.stability_pct,
            round(r.avg_confidence * 100, 1),
            r.frame_count,
            r.cycle_duration,
            sat_counts.get("Rất hài lòng",   0),
            sat_counts.get("Hài lòng",       0),
            sat_counts.get("Bình thường",    0),
            sat_counts.get("Không hài lòng", 0),
        ]

    # ── Truy vấn ─────────────────────────────────────────────────────────────

    def get_all_records(self) -> list[EmotionRecord]:
        """Trả về tất cả bản ghi trong phiên hiện tại."""
        return list(self._records)

    def get_records_as_dicts(self) -> list[dict]:
        """Trả về list dict cho JSON API."""
        return [r.to_display_dict() for r in self._records]

    def get_summary_counts(self) -> dict:
        """Đếm số lần xuất hiện theo từng mức độ hài lòng."""
        counts = {lbl: 0 for lbl in config.SATISFACTION_LEVELS}
        for r in self._records:
            sat = r.satisfaction or config.EMOTION_TO_SATISFACTION_VI.get(r.emotion, r.emotion)
            if sat in counts:
                counts[sat] += 1
            else:
                counts[sat] = 1
        return counts

    def clear(self):
        """Xóa dữ liệu in-memory (bắt đầu phiên mới). Không xóa Sheets/CSV."""
        self._records.clear()
        self._session_id = datetime.now().strftime("%Y%m%d_%H%M%S")
        # Tạo file CSV mới cho phiên mới (nếu dùng CSV)
        if not self._use_sheets:
            self._init_csv_fallback()

    # ── Xuất Excel ───────────────────────────────────────────────────────────

    def export_excel(self, output_path: str = None) -> str:
        """
        Xuất báo cáo Excel (.xlsx) với 2 sheet.

        Parameters
        ----------
        output_path : str | None
            Đường dẫn file đích. Nếu None → tự đặt tên trong EXPORTS_DIR.

        Returns
        -------
        str — đường dẫn file đã xuất.
        """
        try:
            from openpyxl import Workbook
            from openpyxl.styles import (
                Font, PatternFill, Alignment, Border, Side
            )
            from openpyxl.utils import get_column_letter
        except ImportError:
            raise RuntimeError("Cần cài openpyxl: pip install openpyxl")

        if output_path is None:
            fname       = f"EmotionReport_{self._session_id}.xlsx"
            output_path = os.path.join(config.EXPORTS_DIR, fname)

        wb = Workbook()

        # ── Sheet 1: Chi tiết ──────────────────────────────────────────────
        ws1 = wb.active
        ws1.title = "Chi tiết"

        # Style header
        header_fill = PatternFill("solid", fgColor="1E293B")
        header_font = Font(bold=True, color="E2E8F0", size=11)
        center_align = Alignment(horizontal="center", vertical="center")
        thin_border  = Border(
            left=Side(style="thin", color="334155"),
            right=Side(style="thin", color="334155"),
            top=Side(style="thin", color="334155"),
            bottom=Side(style="thin", color="334155"),
        )

        # Headers
        col_headers = [
            "STT", "Mức hài lòng", "Thời gian", "Độ ổn định (%)",
            "Confidence TB (%)", "Số frame", "Thời lượng (s)",
            "Rất hài lòng", "Hài lòng", "Bình thường", "Không hài lòng",
        ]
        ws1.append(col_headers)
        for col_idx, _ in enumerate(col_headers, 1):
            cell = ws1.cell(row=1, column=col_idx)
            cell.fill      = header_fill
            cell.font      = header_font
            cell.alignment = center_align
            cell.border    = thin_border

        # Satisfaction level color mapping cho Excel
        sat_colors = {
            "Rất hài lòng":   "F0FDF4",
            "Hài lòng":        "E0F2FE",
            "Bình thường":    "F8FAFC",
            "Không hài lòng": "FEF2F2",
        }

        # Data rows
        for r in self._records:
            row_data = self._record_to_row(r)
            row_data[4] = f"{row_data[4]}%"   # Confidence dưới dạng %
            row_data[3] = f"{row_data[3]}%"   # Stability
            ws1.append(row_data)

            # Màu row theo mức hài lòng
            sat_label = r.satisfaction or config.EMOTION_TO_SATISFACTION_VI.get(r.emotion, r.emotion)
            bg_color = sat_colors.get(sat_label, "FFFFFF")
            fill = PatternFill("solid", fgColor=bg_color)
            for col_idx in range(1, len(col_headers) + 1):
                cell = ws1.cell(row=ws1.max_row, column=col_idx)
                cell.fill      = fill
                cell.alignment = center_align
                cell.border    = thin_border

        # Điều chỉnh độ rộng cột
        col_widths = [6, 14, 20, 16, 18, 10, 14, 10, 10, 10, 10, 12, 12]
        for i, width in enumerate(col_widths, 1):
            ws1.column_dimensions[get_column_letter(i)].width = width

        ws1.freeze_panes = "A2"

        # ── Sheet 2: Thống kê ──────────────────────────────────────────────
        ws2 = wb.create_sheet("Thống kê")

        ws2.append(["Thống kê tổng hợp phiên demo"])
        ws2.cell(1, 1).font      = Font(bold=True, size=14, color="1E293B")
        ws2.cell(1, 1).alignment = center_align
        ws2.merge_cells("A1:C1")
        ws2.append([])

        ws2.append(["Thông tin phiên"])
        ws2.cell(3, 1).font = Font(bold=True, color="475569")
        ws2.append(["Thời gian bắt đầu", self._session_id.replace("_", " ")])
        ws2.append(["Tổng số bản ghi", len(self._records)])
        ws2.append([])

        # Bảng thống kê theo nhãn
        ws2.append(["Mức độ hài lòng", "Số lượt", "Tỉ lệ (%)"])
        for col_idx in range(1, 4):
            cell = ws2.cell(ws2.max_row, col_idx)
            cell.fill = PatternFill("solid", fgColor="1E293B")
            cell.font = Font(bold=True, color="E2E8F0")
            cell.alignment = center_align

        counts   = self.get_summary_counts()
        total    = sum(counts.values()) or 1
        for sat, count in counts.items():
            pct = count / total * 100
            ws2.append([sat, count, f"{pct:.1f}%"])
            bg = sat_colors.get(sat, "FFFFFF")
            for col_idx in range(1, 4):
                cell = ws2.cell(ws2.max_row, col_idx)
                cell.fill      = PatternFill("solid", fgColor=bg)
                cell.alignment = center_align
                cell.border    = thin_border

        # Tổng
        ws2.append(["TỔNG", len(self._records), "100%"])
        total_row = ws2.max_row
        for col_idx in range(1, 4):
            cell = ws2.cell(total_row, col_idx)
            cell.font   = Font(bold=True, color="1E293B")
            cell.fill   = PatternFill("solid", fgColor="E2E8F0")
            cell.border = thin_border

        for col in ["A", "B", "C"]:
            ws2.column_dimensions[col].width = 22

        wb.save(output_path)
        logger.info(f"Đã xuất Excel: {output_path}")
        return output_path

    # ── Thông tin ─────────────────────────────────────────────────────────────

    @property
    def storage_mode(self) -> str:
        return "Google Sheets" if self._use_sheets else f"CSV ({os.path.basename(self._csv_path or 'local')})"

    @property
    def record_count(self) -> int:
        return len(self._records)
